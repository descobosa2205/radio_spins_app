#!/usr/bin/env python3
"""GENERACIÓN DE INVITACIONES · prueba de regresión (sep 2026).
Comprueba, contra la app REAL y la BD de PRUEBA:
  · el botón «Generar invitaciones» solo sale en lo que promueve una empresa del grupo
  · la pantalla de generación, el PDF de muestra (una página, «INVITACIÓN», QR) y la configuración
    (horas, extras del catálogo y nuevos, imagen, condiciones con plantilla: solo este evento /
    también la plantilla / plantilla nueva)
  · la landing pública de las condiciones (sin sesión) y su token
  · los permisos: los endpoints de dentro resuelven a «invitaciones.gestionar», los públicos a None
  · LOTE 2: las secciones del formato del recinto, generar una categoría (butacas del plano + de pie +
    un sector a mano), el PDF de cada entrada al vuelo y su enlace público, las generadas en la gestión
    de invitaciones (PDF unido, ZIP, plano), recuperar una enviada → código NUEVO (el viejo anulado),
    descartar → anulado, eliminar → anulado, el Excel de códigos y los permisos
  · LOTE 3: la pestaña «Control de accesos» (solo con generadas), el enlace propio (crear, reutilizar,
    anular y generar otro), la página pública sin sesión, las lecturas (OK · ya usada · sin extra ·
    anulada · bloqueada · desconocida · asignada), deshacer, el estado en vivo y las marcas de
    «Invitados», el QR de una autorización de MENORES, el correo y los permisos
    /tmp/python/bin/python3 tools/check_invitaciones_generadas.py
Requiere el entorno de /tmp de CLAUDE.md. Es IDEMPOTENTE (borra lo que crea).
"""
import io
import os
import pathlib
import re
import sys
import tempfile
from datetime import timedelta

RAIZ = pathlib.Path(__file__).resolve().parent.parent
os.chdir(RAIZ)
sys.path.insert(0, str(RAIZ))
tmp = pathlib.Path(tempfile.gettempdir())
for nombre in ("app33_schema_bootstrap.lock", "app33_personnel_bootstrap.lock"):
    (tmp / nombre).write_text("x")
os.environ.setdefault("DATABASE_URL", "postgresql://postgres@127.0.0.1:54329/radiotest?sslmode=disable")
os.environ.setdefault("SUPABASE_URL", "")
os.environ.setdefault("SUPABASE_SERVICE_ROLE_KEY", "")
os.environ.setdefault("FLASK_SECRET_KEY", "check-invgen")
os.environ.setdefault("PGCONNECT_TIMEOUT", "5")

import app as A                                       # noqa: E402
import models                                         # noqa: E402
from models import (Artist, Concert, GroupCompany, InvitationAccessLog, InvitationCategory,   # noqa: E402
                    InvitationConditionsTemplate, InvitationExtraPreset, InvitationGenCategory,
                    InvitationGenConfig, InvitationRequest, InvitationTicket, InvitationVoidedCode,
                    MinorAuthConfig, MinorAuthorization, MinorAuthorizationMinor, Promoter, User,
                    UserProfile, Venue, VenueSeatMap)

A.app.config["WTF_CSRF_ENABLED"] = False
models.ensure_invitation_gen_schema()
OK = FALLOS = 0


def check(nombre, cond, extra=""):
    global OK, FALLOS
    if cond:
        OK += 1; print("  ✓", nombre)
    else:
        FALLOS += 1; print("  ✗", nombre, extra)


A.upload_image = lambda fs, folder, **kw: "https://x.test/storage/invitaciones/entrada.png"
HOY = A.today_local()


def limpia(s):
    for c in s.query(Concert).filter(Concert.festival_name.in_(["InvGen Festival Prueba", "InvGen Tercero Prueba"])).all():
        s.query(InvitationAccessLog).filter(InvitationAccessLog.concert_id == c.id).delete()
        s.query(InvitationVoidedCode).filter(InvitationVoidedCode.concert_id == c.id).delete()
        s.query(InvitationTicket).filter(InvitationTicket.concert_id == c.id).delete()
        for gc in s.query(InvitationGenCategory).filter(InvitationGenCategory.concert_id == c.id).all():
            s.delete(gc)
        s.flush()
        s.query(InvitationCategory).filter(InvitationCategory.concert_id == c.id).delete()
        s.query(InvitationGenConfig).filter(InvitationGenConfig.concert_id == c.id).delete()
        s.delete(c)
    ven = s.query(Venue).filter(Venue.name == "Teatro InvGen").first()
    if ven is not None:
        s.query(VenueSeatMap).filter(VenueSeatMap.venue_id == ven.id).delete()
    for t in s.query(InvitationConditionsTemplate).filter(InvitationConditionsTemplate.name.like("InvGen %")).all():
        s.delete(t)
    for p in s.query(InvitationExtraPreset).filter(InvitationExtraPreset.name.like("InvGen %")).all():
        s.delete(p)
    s.flush()


def datos():
    s = A.db()
    try:
        limpia(s)
        art = s.query(Artist).filter(Artist.name == "Los InvGen").first()
        if art is None:
            art = Artist(name="Los InvGen", photo_url="/static/img/logo.png"); s.add(art); s.flush()
        ven = s.query(Venue).filter(Venue.name == "Teatro InvGen").first()
        if ven is None:
            ven = Venue(name="Teatro InvGen", address="Calle Larga 1", postal_code="11402", municipality="Jerez de la Frontera",
                        province="Cádiz", country="España")
            s.add(ven); s.flush()
        gc = s.query(GroupCompany).filter(GroupCompany.name == "InvGen Producciones").first()
        if gc is None:
            gc = GroupCompany(name="InvGen Producciones", logo_url="/static/img/logo_33_producciones.png"); s.add(gc); s.flush()
        prom = s.query(Promoter).filter(Promoter.nick == "Promotor InvGen").first()
        if prom is None:
            prom = Promoter(nick="Promotor InvGen", contact_email="p@invgen.test"); s.add(prom); s.flush()
        u = s.query(User).filter(User.email == "dir.invgen@prueba.local").first()
        if u is None:
            u = User(email="dir.invgen@prueba.local", password_hash="x", role=10); s.add(u); s.flush()
            s.add(UserProfile(user_id=u.id, nick="dirinvgen", departments=["Ticketing"])); s.flush()
        u.role = 10
        c = Concert(artist_id=art.id, festival_name="InvGen Festival Prueba", activity_type="CONCIERTO",
                    sale_type="VENDIDO", capacity=800, date=HOY + timedelta(days=30), status="CONFIRMADO",
                    venue_id=ven.id, group_company_id=gc.id, billing_company_id=gc.id, doors_time="19:30",
                    show_time="21:00", created_by_user_id=u.id, ticketing_payload={"entry_mode": "SALE"})
        s.add(c)
        c2 = Concert(artist_id=art.id, festival_name="InvGen Tercero Prueba", activity_type="CONCIERTO",
                     sale_type="VENDIDO", capacity=500, date=HOY + timedelta(days=31), status="CONFIRMADO",
                     venue_id=ven.id, promoter_id=prom.id, created_by_user_id=u.id)
        s.add(c2)
        s.commit()
        return {"cid": str(c.id), "cid2": str(c2.id), "uid": str(u.id)}
    finally:
        s.close()


def cliente(uid):
    c = A.app.test_client()
    with c.session_transaction() as ses:
        ses["user_id"] = uid
    return c


def main():
    d = datos()
    cid, cid2, uid = d["cid"], d["cid2"], d["uid"]
    cl = cliente(uid)

    print("1 · el botón y la pantalla")
    r = cl.get(f"/invitaciones/evento/{cid}")
    html = r.get_data(as_text=True)
    check("la pestaña del evento abre", r.status_code == 200, r.status_code)
    check("«Generar invitaciones» sale en la actividad del grupo", "Generar invitaciones" in html)
    r2 = cl.get(f"/invitaciones/evento/{cid2}")
    check("… y NO sale en la de un tercero", r2.status_code == 200 and "Generar invitaciones" not in r2.get_data(as_text=True))
    r3 = cl.get(f"/invitaciones/evento/{cid2}/generar")
    check("la pantalla de un tercero redirige (no se genera)", r3.status_code == 302 and "/invitaciones/evento/" in r3.headers.get("Location", ""))
    r = cl.get(f"/invitaciones/evento/{cid}/generar")
    html = r.get_data(as_text=True)
    check("la pantalla de generación abre", r.status_code == 200, r.status_code)
    check("dice que está sin configurar", "Sin configurar" in html)
    check("el recinto y la dirección salen", "Teatro InvGen" in html and "Calle Larga 1, 11402 Jerez de la Frontera, Cádiz" in html)
    check("las horas de la actividad se ofrecen por defecto", 'value="19:30"' in html and 'value="21:00"' in html)
    check("los tres extras de fábrica están en el catálogo", all(x in html for x in ("Meet &amp; Greet", "After Party", "Parking")))
    check("la plantilla de fábrica se ofrece", "Condiciones generales" in html)
    check("un solo doctype al principio", html.count("<!doctype") == 1 and html.lower().lstrip().startswith("<!doctype"))

    print("2 · el PDF de muestra")
    r = cl.get(f"/invitaciones/evento/{cid}/generar/muestra.pdf")
    pdf = r.get_data()
    check("responde un PDF", r.status_code == 200 and pdf[:4] == b"%PDF", r.status_code)
    from pypdf import PdfReader
    rd = PdfReader(io.BytesIO(pdf))
    txt = "\n".join((p.extract_text() or "") for p in rd.pages)
    check("una sola página", len(rd.pages) == 1, len(rd.pages))
    check("dice INVITACIÓN", "INVITACIÓN" in txt)
    check("lleva el título y el recinto", "Los InvGen" in txt and "Teatro InvGen" in txt)
    check("lleva las horas", "19:30" in txt and "21:00" in txt)
    check("lleva un QR (imagen)", "/Image" in str(rd.pages[0].get("/Resources", {}).get("/XObject", {})) or len(rd.pages[0].images) >= 1)
    check("sin condiciones todavía no promete ninguna", "CONDICIONES DE USO" not in txt)

    print("3 · configurar: horas, extras, imagen y condiciones (plantilla de fábrica modificada → solo este evento)")
    s = A.db()
    try:
        presets = {p.key: p for p in s.query(InvitationExtraPreset).all()}
        tpl = s.query(InvitationConditionsTemplate).filter(InvitationConditionsTemplate.is_builtin.is_(True)).first()
        tpl_id = str(tpl.id)
        tpl_n = len(tpl.clauses_json or [])
        mg_id = str(presets["MEET_GREET"].id)
    finally:
        s.close()
    form = {
        "doors_time": "20:00", "show_time": "21:30",
        "image_choice": "url:https://x.test/cartel.jpg",
        "extra_id[]": ["", ""], "extra_preset_id[]": [mg_id, ""],
        "extra_name[]": ["Meet & Greet", "InvGen Cena previa"], "extra_icon[]": ["fa-handshake", "fa-utensils"],
        "extra_instructions[]": ["A las 19:00 en la puerta de artistas", "Presenta la entrada en el restaurante"],
        "extra_new_catalog[]": ["0", "1"],
        "conditions_template_id": tpl_id, "template_scope": "ONLY",
        "cond_title[]": ["Invitación personal", "Un solo acceso"],
        "cond_body[]": ["Solo este evento la cambia.", "El QR se valida una vez."],
    }
    r = cl.post(f"/invitaciones/evento/{cid}/generar/configurar", data=form, follow_redirects=False)
    check("guarda y vuelve a la pantalla", r.status_code == 302 and "/generar" in r.headers.get("Location", ""), r.status_code)
    s = A.db()
    try:
        cfg = s.query(InvitationGenConfig).filter(InvitationGenConfig.concert_id == A.to_uuid(cid)).first()
        check("hay configuración", cfg is not None)
        check("horas guardadas", cfg.doors_time == "20:00" and cfg.show_time == "21:30")
        check("imagen elegida", cfg.image_url == "https://x.test/cartel.jpg")
        names = [x.name for x in cfg.extras]
        check("dos extras activos, en su orden", names == ["Meet & Greet", "InvGen Cena previa"], names)
        check("el nuevo entra en el catálogo con su icono", any(p.name == "InvGen Cena previa" and p.icon == "fa-utensils" for p in s.query(InvitationExtraPreset).all()))
        check("el extra nuevo cuelga de su preset", all(x.preset_id is not None for x in cfg.extras))
        check("dos cláusulas guardadas en el evento", len(cfg.conditions_json or []) == 2 and cfg.conditions_json[0]["title"] == "Invitación personal")
        tpl = s.get(InvitationConditionsTemplate, A.to_uuid(tpl_id))
        check("«solo este evento»: la plantilla de fábrica NO cambia", len(tpl.clauses_json or []) == tpl_n)
        check("la plantilla queda vinculada", str(cfg.conditions_template_id) == tpl_id)
        check("se creó el token de las condiciones", bool(cfg.conditions_token))
        check("queda configurada y con quién", cfg.configured_at is not None and cfg.configured_by_nick)
        token = cfg.conditions_token
        extra_ids = [str(x.id) for x in cfg.extras]
        extra_presets = [str(x.preset_id) for x in cfg.extras]
    finally:
        s.close()

    print("4 · la pantalla y el PDF con la configuración")
    r = cl.get(f"/invitaciones/evento/{cid}/generar")
    html = r.get_data(as_text=True)
    check("dice Configurados", "Configurados" in html)
    check("enseña los extras como chips", "InvGen Cena previa" in html)
    check("enseña la landing de condiciones", f"/invitaciones/condiciones/{token}" in html)
    r = cl.get(f"/invitaciones/evento/{cid}/generar/muestra.pdf")
    rd = PdfReader(io.BytesIO(r.get_data()))
    txt = "\n".join((p.extract_text() or "") for p in rd.pages)
    check("el PDF lleva las horas configuradas", "20:00" in txt and "21:30" in txt)
    check("el PDF lleva los extras con sus instrucciones", "Cena previa" in txt and "restaurante" in txt)
    check("el PDF resume las condiciones y enlaza a las completas", "CONDICIONES DE USO" in txt and f"/invitaciones/condiciones/{token}" in txt)
    check("sigue siendo una página", len(rd.pages) == 1)
    anots = rd.pages[0].get("/Annots") or []
    check("el enlace de las condiciones es clicable", len(anots) >= 1, len(anots))

    print("5 · la landing pública de las condiciones (sin sesión)")
    anon = A.app.test_client()
    r = anon.get(f"/invitaciones/condiciones/{token}")
    html = r.get_data(as_text=True)
    check("abre sin sesión", r.status_code == 200, r.status_code)
    check("título centrado «Condiciones de uso» y las dos cláusulas", "Condiciones de uso" in html and "Invitación personal" in html and "Un solo acceso" in html)
    check("lleva el logo y el evento", "logo_33_producciones" in html and "Los InvGen" in html)
    check("lleva las og:", 'property="og:title"' in html)
    r = anon.get("/invitaciones/condiciones/no-existe")
    check("un token que no existe da 404 con su página", r.status_code == 404 and "ya no está disponible" in r.get_data(as_text=True))

    print("6 · volver a guardar: cambiar la plantilla TAMBIÉN y crear una nueva")
    form2 = dict(form)
    form2.update({"extra_id[]": extra_ids, "extra_preset_id[]": extra_presets, "extra_new_catalog[]": ["0", "0"],
                  "image_choice": "keep", "template_scope": "UPDATE",
                  "cond_title[]": ["Invitación personal", "Un solo acceso", "Horarios"],
                  "cond_body[]": ["Cambia la plantilla.", "El QR se valida una vez.", "Puertas a la hora indicada."]})
    r = cl.post(f"/invitaciones/evento/{cid}/generar/configurar", data=form2)
    check("segundo guardado", r.status_code == 302)
    s = A.db()
    try:
        cfg = s.query(InvitationGenConfig).filter(InvitationGenConfig.concert_id == A.to_uuid(cid)).first()
        tpl = s.get(InvitationConditionsTemplate, A.to_uuid(tpl_id))
        check("«también la plantilla»: la de fábrica ahora tiene 3 cláusulas", len(tpl.clauses_json or []) == 3 and tpl.clauses_json[0]["body"] == "Cambia la plantilla.")
        check("los extras conservan su id (no se recrean)", [str(x.id) for x in cfg.extras] == extra_ids)
        check("«keep» conserva la imagen", cfg.image_url == "https://x.test/cartel.jpg")
        check("el token de condiciones no cambia", cfg.conditions_token == token)
        # Deshacer la plantilla de fábrica (la prueba es idempotente).
        tpl.clauses_json = [dict(c) for c in A.INVGEN_CONDITIONS_DEFAULT]
        s.commit()
    finally:
        s.close()
    form3 = dict(form2)
    form3.update({"template_scope": "ONLY", "save_template_name": "InvGen Plantilla Nueva", "image_choice": "remove",
                  "extra_id[]": [], "extra_preset_id[]": [], "extra_name[]": [], "extra_icon[]": [], "extra_instructions[]": [], "extra_new_catalog[]": []})
    r = cl.post(f"/invitaciones/evento/{cid}/generar/configurar", data=form3)
    check("tercer guardado", r.status_code == 302)
    s = A.db()
    try:
        cfg = s.query(InvitationGenConfig).filter(InvitationGenConfig.concert_id == A.to_uuid(cid)).first()
        nueva = s.query(InvitationConditionsTemplate).filter(InvitationConditionsTemplate.name == "InvGen Plantilla Nueva").first()
        check("se creó la plantilla nueva con las 3 cláusulas", nueva is not None and len(nueva.clauses_json or []) == 3)
        check("y queda vinculada a la actividad", nueva is not None and cfg.conditions_template_id == nueva.id)
        check("sin extras: se quitan todos", len(cfg.extras) == 0)
        check("«remove» quita la imagen", not cfg.image_url)
        nueva_id = str(nueva.id) if nueva else ""
    finally:
        s.close()
    r = cl.post(f"/invitaciones/evento/{cid}/generar/configurar", data=dict(form3, save_template_name="InvGen Plantilla Nueva"))
    check("un nombre de plantilla repetido se rechaza (302 con el aviso)", r.status_code == 302)
    r = cl.get(f"/invitaciones/evento/{cid}/generar")
    check("… y el aviso reabre el asistente", "invGenConfigModal" in r.get_data(as_text=True) and "Ya hay una plantilla" in r.get_data(as_text=True))
    r = cl.get(f"/invitaciones/generar/plantillas-condiciones/{nueva_id}.json")
    check("el JSON de la plantilla", r.status_code == 200 and r.get_json()["ok"] and len(r.get_json()["clauses"]) == 3)
    r = cl.post(f"/invitaciones/generar/plantillas-condiciones/{nueva_id}/eliminar")
    check("se elimina la plantilla nueva", r.status_code == 200 and r.get_json()["ok"])
    r = cl.post(f"/invitaciones/generar/plantillas-condiciones/{tpl_id}/eliminar")
    check("la de fábrica no se elimina", r.status_code == 400)
    s = A.db()
    try:
        cfg = s.query(InvitationGenConfig).filter(InvitationGenConfig.concert_id == A.to_uuid(cid)).first()
        check("la actividad conserva sus condiciones aunque la plantilla se borre", len(cfg.conditions_json or []) == 3 and cfg.conditions_template_id is None)
    finally:
        s.close()
    r = cl.post(f"/invitaciones/evento/{cid}/generar/configurar", data=dict(form3, **{"cond_title[]": [], "cond_body[]": [], "save_template_name": ""}))
    check("sin cláusulas se rechaza", r.status_code == 302)

    print("7 · permisos")
    with A.app.test_request_context(f"/invitaciones/evento/{cid}/generar"):
        check("la pantalla resuelve a invitaciones.gestionar", A._resolve_request_resource_key() == "invitaciones.gestionar", A._resolve_request_resource_key())
    with A.app.test_request_context(f"/invitaciones/evento/{cid}/generar/configurar", method="POST"):
        check("el guardado resuelve a invitaciones.gestionar", A._resolve_request_resource_key() == "invitaciones.gestionar")
    with A.app.test_request_context(f"/invitaciones/condiciones/{token}"):
        check("la landing pública resuelve a None", A._resolve_request_resource_key() is None)
    check("la landing está en las listas de públicos", "public_invitation_conditions" in A.PUBLIC_ENDPOINTS_EXTRA)
    lector = A.app.test_client()
    with lector.session_transaction() as ses:
        ses["user_id"] = uid
    # Quien NO gestiona invitaciones: un usuario sin rol de dirección ni permisos.
    s = A.db()
    try:
        u2 = s.query(User).filter(User.email == "nadie.invgen@prueba.local").first()
        if u2 is None:
            u2 = User(email="nadie.invgen@prueba.local", password_hash="x", role=1); s.add(u2); s.flush()
            s.add(UserProfile(user_id=u2.id, nick="nadieinvgen", departments=["Diseño"])); s.flush()
        u2.role = 1
        s.commit()
        u2id = str(u2.id)
    finally:
        s.close()
    cl2 = cliente(u2id)
    r = cl2.get(f"/invitaciones/evento/{cid}/generar")
    check("sin acceso a invitaciones no entra (403 o redirección de acceso)", r.status_code in (302, 403), r.status_code)

    # ------------------------------------------------------------------ LOTE 2 · categorías
    print("8 · el formato del recinto y las secciones para la categoría")
    s = A.db()
    try:
        ven = s.query(Venue).filter(Venue.name == "Teatro InvGen").first()
        layout = {"sections": [
            {"id": "s1", "kind": "grid", "name": "Grada", "x": 0, "y": 600, "rows": 3, "cols": 4, "pitch": 26, "rowGap": 30, "rowScheme": "alpha"},
            {"id": "s2", "kind": "floor", "name": "Pista", "x": 0, "y": 250, "w": 400, "h": 200, "cap": 50},
        ], "elements": [
            {"id": "e1", "type": "stage", "x": 0, "y": -100, "w": 300, "h": 120, "label": "ESCENARIO"},
            {"id": "e2", "type": "door", "x": 300, "y": 600, "label": "Puerta 3"},
        ], "categories": []}
        s.add(VenueSeatMap(venue_id=ven.id, name="Principal", is_default=True, layout_json=layout, assignments_json={}, version=1))
        s.commit()
    finally:
        s.close()
    r = cl.get(f"/invitaciones/evento/{cid}/generar/secciones")
    j = r.get_json()
    check("las secciones responden JSON con mapa", r.status_code == 200 and j["ok"] and j["has_map"], r.status_code)
    secs = {x["key"]: x for x in j["sections"]}
    check("Grada numerada con 12 butacas, todas libres", secs["s1"]["numbered"] and secs["s1"]["count"] == 12 and secs["s1"]["free"] == 12, secs.get("s1"))
    check("Pista de pie con aforo 50", not secs["s2"]["numbered"] and secs["s2"]["cap"] == 50)
    check("la puerta del plano se ofrece", "Puerta 3" in j["doors"], j["doors"])
    check("las butacas llevan fila y número impresos", secs["s1"]["seats"][0]["row_label"] == "A" and secs["s1"]["seats"][0]["number"] == "1", secs["s1"]["seats"][:2])
    check("el layout viaja para dibujarlo", bool(j["layout"]["sections"]) and any(e.get("type") == "stage" for e in j["layout"]["elements"]))
    html = cl.get(f"/invitaciones/evento/{cid}/generar").get_data(as_text=True)
    check("la página incluye el asistente de categorías y su botón", "invGenCatModal" in html and "data-invgen-cat-open" in html and "venue_map.js" in html)

    print("9 · generar una categoría: 3 butacas de Grada + 5 de Pista, con puerta")
    seats = secs["s1"]["seats"][:3]
    body = {"name": "InvGen Palco", "extra_ids": [], "sectors": [
        {"section_key": "s1", "section_name": "Grada", "numbered": True, "door": "Puerta 3",
         "seats": [{"key": x["key"], "row_label": x["row_label"], "number": x["number"]} for x in seats], "qty": 0},
        {"section_key": "s2", "section_name": "Pista", "numbered": False, "door": "", "seats": [], "qty": 5},
    ]}
    r = cl.post(f"/invitaciones/evento/{cid}/generar/categorias", json=body)
    j = r.get_json()
    check("se genera (8 invitaciones)", r.status_code == 200 and j["ok"] and j["created"] == 8, (r.status_code, j))
    s = A.db()
    try:
        cat = s.query(InvitationCategory).filter(InvitationCategory.concert_id == A.to_uuid(cid), InvitationCategory.name == "InvGen Palco").first()
        check("hay una categoría de invitación GENERADA, numerada, con 8 por contrato", cat is not None and cat.source == "GENERADA" and cat.ticket_kind == "PDF_NUMBERED" and cat.qty_contract == 8)
        tks = s.query(InvitationTicket).filter(InvitationTicket.category_id == cat.id).all()
        check("8 entradas", len(tks) == 8, len(tks))
        check("todas generadas, con código de 16 y pdf_url NUESTRO", all(len(t.qr_token or "") == 16 and t.is_generated and "/invitaciones/entrada/" in (t.pdf_url or "") and t.ticket_code == t.qr_token for t in tks))
        num = sorted([t for t in tks if t.is_numbered], key=lambda t: (t.row_label, t.seat_number))
        check("3 numeradas con su butaca, fila, map_key y puerta", len(num) == 3 and all(t.sector == "Grada" and t.row_label and t.seat_number and (t.map_key or "").startswith("s1|") and t.door == "Puerta 3" for t in num), [(t.row_label, t.seat_number, t.map_key) for t in num])
        check("5 sin numerar de Pista", sum(1 for t in tks if not t.is_numbered and t.sector == "Pista") == 5)
        check("códigos únicos", len({t.qr_token for t in tks}) == 8)
        gc = s.query(InvitationGenCategory).filter(InvitationGenCategory.concert_id == A.to_uuid(cid)).first()
        check("la categoría generada guarda sus 2 sectores", gc is not None and len(gc.sectors) == 2 and gc.generated_count == 8)
        tok_num, tid_num = num[0].qr_token, str(num[0].id)
        tid_num2 = str(num[1].id)
        tok_qty = [t for t in tks if not t.is_numbered][0].qr_token
        pdf = A._invitation_ticket_pdf_bytes(num[0])
        rd = PdfReader(io.BytesIO(pdf)); txt = "\n".join((p.extract_text() or "") for p in rd.pages)
        check("el PDF de la entrada se compone al vuelo con su butaca, su puerta y su código", pdf[:4] == b"%PDF" and "INVITACIÓN" in txt and "Grada" in txt and "Puerta 3" in txt and tok_num[:4] in txt, txt[:300])
        check("dice la categoría", "InvGen Palco" in txt)
        check("una sola página", len(rd.pages) == 1)
    finally:
        s.close()
    r = cl.get(f"/invitaciones/evento/{cid}/generar/secciones")
    j2 = r.get_json(); secs2 = {x["key"]: x for x in j2["sections"]}
    check("las 3 butacas quedan OCUPADAS para la siguiente categoría", secs2["s1"]["free"] == 9 and all(x["key"] in j2["taken"] for x in seats), secs2["s1"]["free"])
    check("Pista lleva 5 ya generadas", secs2["s2"]["used_unnumbered"] == 5)
    r = cl.post(f"/invitaciones/evento/{cid}/generar/categorias", json=dict(body, name="InvGen Otra"))
    check("repetir esas butacas se rechaza", r.status_code == 400 and "ocupadas" in r.get_json()["error"], r.get_json())
    r = cl.post(f"/invitaciones/evento/{cid}/generar/categorias", json=dict(body, sectors=[body["sectors"][1]]))
    check("un nombre repetido se rechaza", r.status_code == 400 and "Ya hay" in r.get_json()["error"])
    r = cl.post(f"/invitaciones/evento/{cid}/generar/categorias", json={"name": "InvGen Pista grande", "extra_ids": [], "sectors": [{"section_key": "s2", "section_name": "Pista", "numbered": False, "seats": [], "qty": 60}]})
    check("pasarse del aforo se rechaza (quedan 45)", r.status_code == 400 and "aforo" in r.get_json()["error"] and "45" in r.get_json()["error"], r.get_json())
    r = cl.post(f"/invitaciones/evento/{cid}/generar/categorias", json={"name": "InvGen A mano", "extra_ids": [], "sectors": [{"section_key": "", "section_name": "Zona prensa", "numbered": True, "door": "Puerta 1", "seats": [{"key": "", "row_label": "12", "number": "7"}, {"key": "", "row_label": "12", "number": "8"}], "qty": 0}]})
    check("un sector escrito a mano, numerado, se genera", r.status_code == 200 and r.get_json()["created"] == 2, r.get_json())
    r = cl.post(f"/invitaciones/evento/{cid}/generar/categorias", json={"name": "InvGen Vacía", "extra_ids": [], "sectors": []})
    check("sin sectores se rechaza", r.status_code == 400)
    html = cl.get(f"/invitaciones/evento/{cid}/generar").get_data(as_text=True)
    check("la pantalla lista las categorías con sus sectores y el Excel", "InvGen Palco" in html and "Zona prensa" in html and "codigos.xlsx" in html and "8 generadas" in html)

    print("10 · las generadas en la gestión de invitaciones: PDF público, plano, descargas")
    anon = A.app.test_client()
    r = anon.get(f"/invitaciones/entrada/{tok_num}.pdf")
    check("el PDF público abre SIN sesión", r.status_code == 200 and r.get_data()[:4] == b"%PDF", r.status_code)
    check("… y no se cachea", "no-store" in r.headers.get("Cache-Control", ""))
    check("el código con guiones o en minúsculas también vale", anon.get(f"/invitaciones/entrada/{tok_num[:4].lower()}-{tok_num[4:]}.pdf").status_code == 200)
    check("un código que no existe da 404", anon.get("/invitaciones/entrada/NOEXISTE12345678.pdf").status_code == 404)
    r = cl.get(f"/invitaciones/evento/{cid}")
    html = r.get_data(as_text=True)
    check("la gestión del evento enseña la categoría generada con el plano de su sector", r.status_code == 200 and "InvGen Palco" in html and 'data-sector="Grada"' in html)
    # Las entradas viajan en el parcial del asignador (la página no las incrusta: se piden al abrirlo).
    r = cl.get(f"/invitaciones/evento/{cid}/asignador-parcial")
    parcial = r.get_data(as_text=True)
    check("el asignador lleva la entrada generada (por su id)", r.status_code == 200 and tid_num in parcial, r.status_code)
    check("… y su butaca del plano (map_key) ya resuelta", "s1|1|1" in parcial or "s1|1|2" in parcial, parcial.count("map_key"))
    s = A.db()
    try:
        tks = s.query(InvitationTicket).filter(InvitationTicket.concert_id == A.to_uuid(cid), InvitationTicket.is_generated.is_(True)).order_by(InvitationTicket.uploaded_at).all()
        pdfm, n = A._invitation_tickets_to_merged_pdf(tks[:3])
        check("el PDF unido de 3 generadas tiene 3 páginas (compuestas al vuelo)", n == 3 and len(PdfReader(io.BytesIO(pdfm)).pages) == 3, n)
        zb, zname, added = A._invitation_tickets_to_zip(tks[:2])
        check("el ZIP incluye las 2", added == 2 and zb[:2] == b"PK")
    finally:
        s.close()
    r = cl.post(f"/invitaciones/evento/{cid}/tickets/redetectar", data={"ajax": "1"})
    check("re-detectar butacas se salta las generadas", r.status_code == 200 and r.get_json()["total"] == 0, r.get_json() if r.status_code == 200 else r.status_code)

    print("11 · recuperar una enviada → código NUEVO; descartar → anulado; eliminar → anulado")
    s = A.db()
    try:
        t = s.get(InvitationTicket, A.to_uuid(tid_num))
        t.status = "SENT"; t.sent_at = A._now_madrid(); t.assigned_label = "Pepe Invitado"
        t2 = s.get(InvitationTicket, A.to_uuid(tid_num2))
        t2.status = "SENT"; t2.sent_at = A._now_madrid(); t2.assigned_label = "Otro Invitado"
        s.commit()
    finally:
        s.close()
    r = cl.post(f"/invitaciones/tickets/{tid_num}/liberar", data={"ajax": "1"})
    check("una enviada pide confirmación", r.status_code == 200 and r.get_json().get("needs_confirm"))
    r = cl.post(f"/invitaciones/tickets/{tid_num}/liberar", data={"ajax": "1", "mode": "recover"})
    check("se recupera y lo dice: código nuevo", r.status_code == 200 and r.get_json()["ok"] and "código nuevo" in r.get_json()["message"], r.get_json())
    s = A.db()
    try:
        t = s.get(InvitationTicket, A.to_uuid(tid_num))
        check("la entrada vuelve a disponible con OTRO código y versión 2", t.status == "AVAILABLE" and t.qr_token != tok_num and len(t.qr_token) == 16 and t.code_version == 2 and t.ticket_code == t.qr_token)
        check("su pdf_url apunta al código nuevo", t.qr_token in (t.pdf_url or ""))
        check("el código viejo queda ANULADO", s.query(InvitationVoidedCode).filter(InvitationVoidedCode.qr_token == tok_num).count() == 1)
        check("su map_key y su butaca se conservan", (t.map_key or "").startswith("s1|") and t.sector == "Grada")
        tok_new = t.qr_token
    finally:
        s.close()
    check("el PDF del código viejo ya no se sirve (410)", anon.get(f"/invitaciones/entrada/{tok_num}.pdf").status_code == 410)
    check("… y el nuevo sí", anon.get(f"/invitaciones/entrada/{tok_new}.pdf").status_code == 200)
    r = cl.post(f"/invitaciones/tickets/{tid_num2}/liberar", data={"ajax": "1", "mode": "lost"})
    check("descartar una enviada", r.status_code == 200 and r.get_json()["ok"])
    s = A.db()
    try:
        t2 = s.get(InvitationTicket, A.to_uuid(tid_num2))
        check("queda LOST con su código anulado y sin renacer", t2.status == "LOST" and t2.code_version == 1 and s.query(InvitationVoidedCode).filter(InvitationVoidedCode.qr_token == t2.qr_token).count() == 1)
        tok_lost = t2.qr_token
        # el recuperar EN BLOQUE (compromiso/solicitud) pasa por _invitation_release_apply
        t3 = [x for x in s.query(InvitationTicket).filter(InvitationTicket.concert_id == A.to_uuid(cid), InvitationTicket.is_generated.is_(True), InvitationTicket.status == "AVAILABLE").all() if not x.is_numbered][0]
        t3.status = "SENT"; t3.sent_at = A._now_madrid(); s.flush()
        tok3 = t3.qr_token
        rec, disc = A._invitation_release_apply([t3], mode="recover", entity_label="la prueba")
        s.commit()
        check("el recuperar en bloque también renace el código", rec == 1 and t3.qr_token != tok3 and s.query(InvitationVoidedCode).filter(InvitationVoidedCode.qr_token == tok3).count() == 1)
        # una NO enviada que se libera no cambia de código (nadie lo tenía)
        t4 = [x for x in s.query(InvitationTicket).filter(InvitationTicket.concert_id == A.to_uuid(cid), InvitationTicket.is_generated.is_(True), InvitationTicket.status == "AVAILABLE").all() if not x.is_numbered and x.id != t3.id][0]
        t4.status = "ASSIGNED"; t4.assigned_label = "Alguien"; s.flush(); tok4 = t4.qr_token
        A._invitation_release_apply([t4], mode="recover", entity_label="la prueba"); s.commit()
        check("una asignada sin enviar conserva su código al liberarla", t4.qr_token == tok4 and t4.status == "AVAILABLE")
        tid4 = str(t4.id)
    finally:
        s.close()
    check("el PDF de la descartada da 410", anon.get(f"/invitaciones/entrada/{tok_lost}.pdf").status_code == 410)
    r = cl.post(f"/invitaciones/tickets/{tid_num}/actualizar", data={"ticket_code": "OTRO", "sector": "X"}, follow_redirects=False)
    s = A.db()
    try:
        t = s.get(InvitationTicket, A.to_uuid(tid_num))
        check("una generada no se edita a mano (se rebota sin tocarla)", r.status_code == 302 and t.ticket_code == tok_new and t.sector == "Grada")
    finally:
        s.close()
    r = cl.post(f"/invitaciones/tickets/{tid4}/eliminar", follow_redirects=False)
    s = A.db()
    try:
        check("eliminar una generada la borra y anula su código", r.status_code == 302 and s.get(InvitationTicket, A.to_uuid(tid4)) is None and s.query(InvitationVoidedCode).filter(InvitationVoidedCode.qr_token == tok4).count() == 1)
    finally:
        s.close()

    print("12 · el Excel de códigos y eliminar la categoría")
    r = cl.get(f"/invitaciones/evento/{cid}/generar/codigos.xlsx")
    check("el Excel se descarga", r.status_code == 200 and r.get_data()[:2] == b"PK" and "xlsx" in r.headers.get("Content-Disposition", ""), r.status_code)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.get_data()))
    ws = wb["Códigos"]; ws2 = wb["Anulados"]
    codes = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
    check("lleva los códigos vivos y la hoja de anulados con los viejos", tok_new in codes and tok_num not in codes and tok_num in [row[0] for row in ws2.iter_rows(min_row=2, values_only=True)])
    s = A.db()
    try:
        gc = s.query(InvitationGenCategory).filter(InvitationGenCategory.concert_id == A.to_uuid(cid), InvitationGenCategory.name == "InvGen Palco").first()
        gid = str(gc.id)
        t = s.get(InvitationTicket, A.to_uuid(tid_num)); t.status = "ASSIGNED"; t.assigned_label = "Pepe"; s.commit()
        vivos = s.query(InvitationTicket).filter(InvitationTicket.gen_category_id == gc.id).count()
    finally:
        s.close()
    r = cl.post(f"/invitaciones/evento/{cid}/generar/categorias/{gid}/eliminar", follow_redirects=False)
    s = A.db()
    try:
        check("con una asignada NO se elimina", r.status_code == 302 and s.query(InvitationGenCategory).get(A.to_uuid(gid)) is not None)
        t = s.get(InvitationTicket, A.to_uuid(tid_num)); t.status = "AVAILABLE"; t.assigned_label = None; s.commit()
        n_void_antes = s.query(InvitationVoidedCode).filter(InvitationVoidedCode.concert_id == A.to_uuid(cid)).count()
    finally:
        s.close()
    r = cl.post(f"/invitaciones/evento/{cid}/generar/categorias/{gid}/eliminar", follow_redirects=False)
    s = A.db()
    try:
        check("sin asignadas se elimina con sus entradas y su categoría", r.status_code == 302 and s.get(InvitationGenCategory, A.to_uuid(gid)) is None and s.query(InvitationTicket).filter(InvitationTicket.gen_category_id == A.to_uuid(gid)).count() == 0 and s.query(InvitationCategory).filter(InvitationCategory.name == "InvGen Palco").count() == 0)
        n_void = s.query(InvitationVoidedCode).filter(InvitationVoidedCode.concert_id == A.to_uuid(cid)).count()
        check("y todos sus códigos quedan anulados", n_void == n_void_antes + vivos, (n_void, n_void_antes, vivos))
    finally:
        s.close()

    print("13 · permisos del lote 2")
    with A.app.test_request_context(f"/invitaciones/evento/{cid}/generar/secciones"):
        check("las secciones resuelven a invitaciones.gestionar", A._resolve_request_resource_key() == "invitaciones.gestionar")
    with A.app.test_request_context(f"/invitaciones/evento/{cid}/generar/categorias", method="POST"):
        check("generar resuelve a invitaciones.gestionar", A._resolve_request_resource_key() == "invitaciones.gestionar")
    with A.app.test_request_context(f"/invitaciones/entrada/{tok_new}.pdf"):
        check("el PDF público resuelve a None", A._resolve_request_resource_key() is None)
    check("el PDF público está en las listas de públicos", "public_invitation_ticket_pdf" in A.PUBLIC_ENDPOINTS_EXTRA)
    check("sin acceso: las secciones no se ven", cl2.get(f"/invitaciones/evento/{cid}/generar/secciones").status_code in (302, 403))


    # ------------------------------------------------------------------ LOTE 3 · control de acceso
    print("14 · la pestaña «Control de accesos» y el enlace propio")
    # Volver a activar dos extras (el apartado 6 los quitó todos) y generar una categoría CON extras.
    s = A.db()
    try:
        presets = {p.key: p for p in s.query(InvitationExtraPreset).all()}
        cena = s.query(InvitationExtraPreset).filter(InvitationExtraPreset.name == "InvGen Cena previa").first()
        mg_id, cena_id = str(presets["MEET_GREET"].id), (str(cena.id) if cena else "")
        cfg = s.query(InvitationGenConfig).filter(InvitationGenConfig.concert_id == A.to_uuid(cid)).first()
        clauses = list(cfg.conditions_json or [])
    finally:
        s.close()
    r = cl.post(f"/invitaciones/evento/{cid}/generar/configurar", data={
        "doors_time": "20:00", "show_time": "21:30", "image_choice": "keep",
        "extra_id[]": ["", ""], "extra_preset_id[]": [mg_id, cena_id],
        "extra_name[]": ["Meet & Greet", "InvGen Cena previa"], "extra_icon[]": ["fa-handshake", "fa-utensils"],
        "extra_instructions[]": ["A las 19:00", "En el restaurante"], "extra_new_catalog[]": ["0", "0"],
        "template_scope": "ONLY",
        "cond_title[]": [c["title"] for c in clauses] or ["Invitación personal"],
        "cond_body[]": [c["body"] for c in clauses] or ["Solo este evento."],
    })
    check("se vuelven a activar los dos extras", r.status_code == 302)
    s = A.db()
    try:
        cfg = s.query(InvitationGenConfig).filter(InvitationGenConfig.concert_id == A.to_uuid(cid)).first()
        extras_now = {x.name: str(x.id) for x in cfg.extras}
        check("hay dos extras activos", set(extras_now) == {"Meet & Greet", "InvGen Cena previa"}, extras_now)
        mg_key = extras_now["Meet & Greet"]
        cena_key = extras_now["InvGen Cena previa"]
    finally:
        s.close()
    r = cl.post(f"/invitaciones/evento/{cid}/generar/categorias", json={"name": "InvGen Acceso", "extra_ids": [mg_key], "sectors": [
        {"section_key": "s2", "section_name": "Pista", "numbered": False, "door": "Puerta 2", "seats": [], "qty": 3}]})
    check("se genera una categoría con el extra Meet & Greet (3 de pie)", r.status_code == 200 and r.get_json()["created"] == 3, r.get_json())
    s = A.db()
    try:
        vivas = (s.query(InvitationTicket).filter(InvitationTicket.concert_id == A.to_uuid(cid), InvitationTicket.is_generated.is_(True),
                                                  InvitationTicket.status != "LOST").order_by(InvitationTicket.uploaded_at.asc()).all())
        con_mg = [t for t in vivas if t.gen_category is not None and mg_key in {str(x) for x in (t.gen_category.extras_json or [])}]
        sin_mg = [t for t in vivas if t not in con_mg]
        check("hay entradas con el extra y sin él", len(con_mg) == 3 and len(sin_mg) >= 1, (len(con_mg), len(sin_mg)))
        n_validas = len([t for t in vivas if (t.status or "").upper() != "BLOCKED"])
        tk_a, tk_b, tk_c = con_mg[0], con_mg[1], con_mg[2]
        tk_sin = sin_mg[0]
        cod_a, cod_b, cod_c, cod_sin = tk_a.qr_token, tk_b.qr_token, tk_c.qr_token, tk_sin.qr_token
        tid_a, tid_b, tid_c = str(tk_a.id), str(tk_b.id), str(tk_c.id)
    finally:
        s.close()
    r = cl.get(f"/invitaciones/evento/{cid}")
    html = r.get_data(as_text=True)
    check("la pestaña sale en la gestión del evento", r.status_code == 200 and 'id="inv-tab-access"' in html and "Control de accesos" in html)
    check("con sus dos botones", "Compartir códigos de entradas" in html and "Control de acceso propio" in html and "codigos.xlsx" in html)
    check("una tarjeta por control: la entrada y los dos extras", html.count("data-invacc-card=") == 3 and "Acceso entrada" in html and "InvGen Cena previa" in html)
    check("las marcas de Invitados van preparadas (aunque no haya nada) ", "invgen_access.js" in html)
    r2 = cl.get(f"/invitaciones/evento/{cid2}")
    check("… y NO sale en la actividad sin generadas", r2.status_code == 200 and 'id="inv-tab-access"' not in r2.get_data(as_text=True))
    r = cl.post(f"/invitaciones/evento/{cid}/generar/control/enlace")
    j = r.get_json()
    check("se genera el enlace", r.status_code == 200 and j["ok"] and "/control-acceso/" in j["url"], j)
    acc_url = j["url"]
    tok_acc = acc_url.rstrip("/").rsplit("/", 1)[-1]
    check("el título de lo que se comparte", j["share"]["title"].startswith("Control de accesos · Concierto · Los InvGen"), j["share"]["title"])
    check("y debajo el festival y la fecha", "InvGen Festival Prueba" in j["share"]["description"] and HOY.strftime("%Y") in j["share"]["description"], j["share"]["description"])
    check("WhatsApp y SMS con el enlace dentro", j["share"]["whatsapp_url"].startswith("https://wa.me/") and "control-acceso" in j["share"]["sms_url"])
    r = cl.post(f"/invitaciones/evento/{cid}/generar/control/enlace")
    check("volver a pedirlo devuelve el MISMO enlace", r.get_json()["url"] == acc_url)
    r = cl.post(f"/invitaciones/evento/{cid}/generar/control/enlace", data={"renew": "1"})
    acc_url2 = r.get_json()["url"]
    check("«anular y generar otro» cambia el enlace", acc_url2 != acc_url and r.get_json()["renewed"])
    anon = A.app.test_client()
    check("el enlace anterior ya no vale (404)", anon.get(acc_url).status_code == 404)
    acc_url, tok_acc = acc_url2, acc_url2.rstrip("/").rsplit("/", 1)[-1]
    r = cl.get(f"/invitaciones/evento/{cid}/generar/control/estado.json")
    j = r.get_json()
    check("el estado responde", r.status_code == 200 and j["ok"] and j["controls"][0]["key"] == "ENTRADA" and len(j["controls"]) == 3, r.status_code)
    ent = j["controls"][0]
    check("emitidas de la entrada = las generadas válidas", ent["issued"] == n_validas and ent["validated"] == 0, (ent, n_validas))
    mgc = [c for c in j["controls"] if c["key"] == mg_key][0]
    check("el Meet & Greet cuenta solo las que lo incluyen", mgc["issued"] == 3, mgc)
    check("el estado sin personas no las trae", j.get("people") is None)
    r = cl.get(f"/invitaciones/evento/{cid}/generar/control/estado.json?people=1")
    j = r.get_json()
    check("con people=1 trae una fila por entrada, sin asignar", len(j["people"]) >= 4 and all(p["guest_name"] == "Sin asignar" for p in j["people"] if p["valid"]))

    print("15 · la página pública y las lecturas")
    r = anon.get(acc_url)
    html = r.get_data(as_text=True)
    check("la página del control abre SIN sesión", r.status_code == 200, r.status_code)
    check("lleva las og: con el título y la imagen", 'og:title" content="Control de accesos' in html and "/og.jpg" in html)
    check("enseña la cabecera y los controles", "Los InvGen" in html and "Acceso entrada" in html and "InvGen Cena previa" in html and "data-scan-url" in html)
    check("todavía sin menores (el atributo no se emite)", 'data-minors-check-url="' not in html)
    check("un solo doctype", html.count("<!doctype") == 1)
    r = anon.get(f"/control-acceso/{tok_acc}/og.jpg")
    check("la miniatura responde (imagen o el respaldo)", r.status_code in (200, 302), r.status_code)
    r = anon.get(f"/control-acceso/{tok_acc}/estado.json")
    check("el estado público responde solo los números", r.status_code == 200 and r.get_json()["ok"] and "people" not in r.get_json() and "by_source" not in r.get_json())

    def leer(code, control="ENTRADA"):
        rr = anon.post(f"/control-acceso/{tok_acc}/leer", json={"code": code, "control": control})
        return rr.status_code, (rr.get_json() or {})

    st, j = leer(cod_a)
    check("una entrada válida → OK", st == 200 and j["ok"] and j["result"] == "OK" and j["good"], j)
    check("… dice que no está asignada a nadie", "no está asignada" in (j.get("warn") or ""), j.get("warn"))
    check("… y el estado que vuelve ya cuenta 1 dentro", j["state"]["controls"][0]["validated"] == 1)
    st, j = leer(cod_a)
    check("la misma entrada otra vez → YA USADA (rojo)", j["result"] == "YA_USADA" and not j["good"] and "Ya entró" in j["detail"], j)
    st, j = leer(cod_a, mg_key)
    check("el Meet & Greet de una entrada que lo incluye → OK", j["result"] == "OK" and j["good"] and "Meet & Greet OK" in j["title"], j)
    st, j = leer(cod_a, mg_key)
    check("el mismo extra otra vez → YA USADA", j["result"] == "YA_USADA" and "Se usó" in j["detail"], j)
    st, j = leer(cod_a, cena_key)
    check("un extra que la entrada NO incluye → SIN EXTRA", j["result"] == "SIN_EXTRA" and not j["good"], j)
    st, j = leer(cod_sin, mg_key)
    check("una entrada de otra categoría sin el extra → SIN EXTRA", j["result"] == "SIN_EXTRA", j)
    st, j = leer(f"https://app.33producciones.es/invitaciones/entrada/{cod_b[:4].lower()}-{cod_b[4:]}.pdf")
    check("el código llega como la URL del PDF (con guiones y minúsculas) y se lee igual", j["result"] == "OK" and j["ticket"]["code"] == cod_b, j)
    st, j = leer(tok_num)
    check("el código ANULADO del apartado 11 → ANULADA", j["result"] == "ANULADA" and not j["good"], j)
    st, j = leer("NOEXISTE12345678")
    check("un código que no existe → DESCONOCIDA", j["result"] == "DESCONOCIDA" and not j["good"], j)
    st, j = leer("")
    check("sin código, 400", st == 400 and not j.get("ok"))
    s = A.db()
    try:
        t = s.get(InvitationTicket, A.to_uuid(tid_c)); t.status = "BLOCKED"; s.commit()
    finally:
        s.close()
    st, j = leer(cod_c)
    check("una entrada BLOQUEADA → BLOQUEADA (rojo)", j["result"] == "BLOQUEADA" and not j["good"], j)
    s = A.db()
    try:
        t = s.get(InvitationTicket, A.to_uuid(tid_c)); t.status = "ASSIGNED"; t.assigned_label = "Pepe Puerta"; s.commit()
        a = s.get(InvitationTicket, A.to_uuid(tid_a))
        check("la entrada leída queda con su hora de acceso y su extra usado", a.access_entered_at is not None and mg_key in (a.access_extras_json or {}))
        n_logs = s.query(InvitationAccessLog).filter(InvitationAccessLog.concert_id == A.to_uuid(cid)).count()
        res = {x.result for x in s.query(InvitationAccessLog).filter(InvitationAccessLog.concert_id == A.to_uuid(cid)).all()}
        check("cada lectura deja su traza (también los rechazos)", n_logs >= 9 and {"OK", "YA_USADA", "SIN_EXTRA", "ANULADA", "DESCONOCIDA", "BLOQUEADA"} <= res, (n_logs, res))
    finally:
        s.close()
    st, j = leer(cod_c)
    check("una entrada asignada dice de quién es y sin aviso", j["result"] == "OK" and j["ticket"]["guest_name"] == "Pepe Puerta" and not j.get("warn"), j)
    r = cl.post(f"/invitaciones/evento/{cid}/generar/control/deshacer", json={"ticket_id": tid_c, "control": "ENTRADA"})
    j = r.get_json()
    check("deshacer una lectura (desde dentro)", r.status_code == 200 and j["ok"] and j["undone"] and j["state"]["controls"][0]["validated"] == 2, j.get("state", {}).get("controls", [{}])[0])
    s = A.db()
    try:
        c3 = s.get(InvitationTicket, A.to_uuid(tid_c))
        check("… y la entrada vuelve a estar sin pasar", c3.access_entered_at is None)
        check("… con su traza DESHECHO", s.query(InvitationAccessLog).filter(InvitationAccessLog.ticket_id == c3.id, InvitationAccessLog.result == "DESHECHO").count() == 1)
        # recuperar una enviada: renace con otro código y SIN usos
        a = s.get(InvitationTicket, A.to_uuid(tid_a)); a.status = "SENT"; a.sent_at = A._now_madrid(); s.commit()
    finally:
        s.close()
    r = cl.post(f"/invitaciones/tickets/{tid_a}/liberar", data={"ajax": "1", "mode": "recover"})
    s = A.db()
    try:
        a = s.get(InvitationTicket, A.to_uuid(tid_a))
        check("al recuperarla renace sin accesos ni extras usados", r.status_code == 200 and a.qr_token != cod_a and a.access_entered_at is None and (a.access_extras_json or {}) == {})
        cod_a_nuevo = a.qr_token
    finally:
        s.close()
    st, j = leer(cod_a)
    check("y su código viejo ya dice ANULADA en la puerta", j["result"] == "ANULADA", j)
    st, j = leer(cod_a_nuevo)
    check("el nuevo entra", j["result"] == "OK", j)

    print("16 · lo que se ve: las listas de la pestaña y las marcas de «Invitados»")
    s = A.db()
    try:
        cat = s.query(InvitationCategory).filter(InvitationCategory.concert_id == A.to_uuid(cid), InvitationCategory.name == "InvGen Acceso").first()
        req = InvitationRequest(concert_id=A.to_uuid(cid), guest_name="Invitada Acceso", guest_type="MANUAL", requester_nick="dirinvgen",
                                quantities_json={str(cat.id): 1}, status="ASIGNADAS")
        s.add(req); s.flush()
        b = s.get(InvitationTicket, A.to_uuid(tid_b)); b.status = "ASSIGNED"; b.assigned_request_id = req.id; b.assigned_label = "Invitada Acceso"
        s.commit()
        req_id = str(req.id)
    finally:
        s.close()
    r = cl.get(f"/invitaciones/evento/{cid}/generar/control/estado.json?people=1")
    j = r.get_json()
    fila = [p for p in j["people"] if p["id"] == tid_b][0]
    check("la fila de la entrada dice de quién es y que ya ha entrado", fila["guest_name"] == "Invitada Acceso" and fila["entered"] and fila["source"] == f"request:{req_id}", fila)
    check("por petición: 1 de 1 dentro", j["by_source"].get(f"request:{req_id}", {}).get("entered") == 1 and j["by_source"][f"request:{req_id}"]["total"] == 1, j["by_source"])
    s = A.db()
    try:
        c3 = s.get(InvitationTicket, A.to_uuid(tid_c)); c3.status = "BLOCKED"; s.commit()   # una apartada, para verla en la lista oculta
    finally:
        s.close()
    r = cl.get(f"/invitaciones/evento/{cid}")
    html = r.get_data(as_text=True)
    s = A.db()
    try:
        c3 = s.get(InvitationTicket, A.to_uuid(tid_c)); c3.status = "ASSIGNED"; s.commit()
    finally:
        s.close()
    check("la pestaña pinta la fila en verde (is-on) con su nombre", f'data-invacc-row="{tid_b}" data-control="ENTRADA"' in html and "Invitada Acceso" in html and re.search(r'class="invacc-row is-on"[^>]*data-invacc-row="%s" data-control="ENTRADA"' % tid_b, html) is not None)
    check("y en «Invitados» sale la marca «1/1 dentro» de esa petición", f'data-access-src="request:{req_id}"' in html and re.search(r'data-access-src="request:%s"(?:(?!</span>\s*</span>).)*?1</span>/<span data-access-total-n>1</span> dentro' % req_id, html, re.S) is not None)
    check("el número de la pestaña dice cuántos han entrado", re.search(r'data-invacc-pill[^>]*>\s*2\s*<', html) is not None)
    check("la bloqueada va en una fila oculta con su botón de verlas", "Ver también las anuladas y bloqueadas (1)" in html and f'data-invacc-row="{tid_c}" data-control="ENTRADA" data-acc="BLOCKED"' in html)
    check("la pestaña carga su JS", "invgen_access.js" in html)

    print("17 · menores: el mismo lector entiende el QR de una autorización")
    s = A.db()
    try:
        mcfg = MinorAuthConfig(concert_id=A.to_uuid(cid), age_limit=16, public_token=A._uuid_token(), validate_token=A._uuid_token(), active=True)
        s.add(mcfg); s.flush()
        auth = MinorAuthorization(config_id=mcfg.id, concert_id=A.to_uuid(cid), guardian_kind="MADRE", guardian_first_name="Ana", guardian_last_name="Tutora",
                                  guardian_doc_number="12345678Z", escort_is_guardian=True, qr_token=A._uuid_token(), status="VALID")
        s.add(auth); s.flush()
        s.add(MinorAuthorizationMinor(authorization_id=auth.id, first_name="Leo", last_name="Menor", doc_number="87654321X"))
        auth2 = MinorAuthorization(config_id=mcfg.id, concert_id=A.to_uuid(cid), guardian_kind="TUTOR", guardian_first_name="Luis", guardian_last_name="Anulado",
                                   qr_token=A._uuid_token(), status="CANCELLED")
        s.add(auth2); s.commit()
        qr_ok, qr_ko = auth.qr_token, auth2.qr_token
    finally:
        s.close()
    st, j = leer(f"https://app.33producciones.es/autorizacion-menores/pase/{qr_ok}")
    check("el QR de una autorización válida → MENOR OK (verde)", j["result"] == "MENOR_OK" and j["good"] and j["kind"] == "minor" and j["minor"]["minors"][0]["full_name"] == "Leo Menor", j)
    st, j = leer(qr_ko)
    check("el de una anulada → MENOR KO (rojo)", j["result"] == "MENOR_KO" and not j["good"], j)
    r = anon.get(acc_url)
    html = r.get_data(as_text=True)
    check("la página ofrece ahora el acceso de menores con el buscador de siempre", "Acceso de menores" in html and 'data-minors-check-url="' in html and "doc_camera.js" in html)
    check("… y las entradas se siguen leyendo igual", leer(cod_sin)[1]["result"] == "OK")

    print("18 · el correo y los permisos del lote 3")
    enviados = []
    _orig_send = A._send_optional_email
    A._send_optional_email = lambda to, subject, html, *a, **k: (enviados.append((to, subject, html)) or (True, None))
    try:
        r = cl.post(f"/invitaciones/evento/{cid}/generar/control/enviar", json={"emails": "puerta@x.test, seguridad@x.test", "note": "Nota de prueba"})
        j = r.get_json()
        check("el correo sale a las dos personas", r.status_code == 200 and j["ok"] and j["sent"] == 2 and len(enviados) == 1 and set(enviados[0][0]) == {"puerta@x.test", "seguridad@x.test"}, j)
        to, subject, html = enviados[0] if enviados else ([], "", "")
        check("con el asunto de la casa", subject.startswith("Control de accesos · Concierto · Los InvGen"), subject)
        check("logo a la derecha, título centrado, la actividad y el botón con el enlace", "text-align:right" in html and "Control de acceso</h2>" in html and "Los InvGen" in html and acc_url in html and "Nota de prueba" in html)
        r = cl.post(f"/invitaciones/evento/{cid}/generar/control/enviar", json={"emails": "", "note": ""})
        check("sin correos se rechaza", r.status_code == 400)
    finally:
        A._send_optional_email = _orig_send
    with A.app.test_request_context(f"/invitaciones/evento/{cid}/generar/control/estado.json"):
        check("el estado de dentro resuelve a invitaciones.gestionar", A._resolve_request_resource_key() == "invitaciones.gestionar")
    with A.app.test_request_context(f"/invitaciones/evento/{cid}/generar/control/enlace", method="POST"):
        check("el enlace resuelve a invitaciones.gestionar", A._resolve_request_resource_key() == "invitaciones.gestionar")
    with A.app.test_request_context(f"/control-acceso/{tok_acc}"):
        check("la página pública resuelve a None", A._resolve_request_resource_key() is None)
    with A.app.test_request_context(f"/control-acceso/{tok_acc}/leer", method="POST"):
        check("la lectura pública resuelve a None", A._resolve_request_resource_key() is None)
    check("los cuatro públicos están en las listas", all(e in A.PUBLIC_ENDPOINTS_EXTRA for e in ("public_invitation_access", "public_invitation_access_state", "public_invitation_access_scan", "public_invitation_access_og_image")))
    vf = A.app.view_functions["public_invitation_access_scan"]
    check("la lectura pública está EXENTA de CSRF de verdad", ("%s.%s" % (vf.__module__, vf.__name__)) in A.csrf._exempt_views)
    check("sin acceso: el estado de dentro no se ve", cl2.get(f"/invitaciones/evento/{cid}/generar/control/estado.json").status_code in (302, 403))
    check("sin acceso: el enlace no se genera", cl2.post(f"/invitaciones/evento/{cid}/generar/control/enlace").status_code in (302, 403))
    r = anon.get(f"/invitaciones/evento/{cid}")
    check("la gestión del evento sigue exigiendo sesión", r.status_code in (302, 403))

    print()
    print(f"OK: {OK} · FALLOS: {FALLOS}")
    return 0 if FALLOS == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
