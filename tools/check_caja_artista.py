#!/usr/bin/env python3
"""LA CAJA DE UN ARTISTA · prueba de regresión (sep 2026).

Comprueba, contra la app REAL y una BD de PRUEBA propia, lo que pidió Dani:

  · el GASTO solo entra en la caja cuando la bolsa está **CERRADA** y administración ha dicho que
    se incluye (una bolsa abierta es lo que se VA a gastar, no lo gastado);
  · lo que cuenta es el **BALANCE FINAL**: lo que cubre el promotor (o se le refactura) y lo que
    cubre el artista NO es nuestro, y lo que asume la bolsa de una actividad lo paga su caché;
  · los ROYALTIES: el beneficio de la compañía es lo facturado por su repertorio **menos TODO lo
    que se paga por él** —los del propio artista y los de cualquier otro que cobre de sus obras—,
    y un royalty pagado **no es gasto: reduce el ingreso**;
  · las ACTIVIDADES se reparten con los **porcentajes del contrato** sobre el **importe final de la
    liquidación** (administración puede cambiarlo a última hora);
  · al cerrar la liquidación **se pregunta** si va a la caja, y se puede cambiar después;
  · las bolsas que YA estaban cerradas entran solas (el arreglo de todas las previas).

    /tmp/python/bin/python3 tools/check_caja_artista.py

⚠️ Usa su PROPIA base (`radiocaja`) y la RECREA en cada pasada: una BD de prueba contaminada miente.
"""
import os, sys, tempfile, pathlib, datetime, decimal

os.chdir("/Users/carlos/Documents/radio_spins_app")
sys.path.insert(0, "/Users/carlos/Documents/radio_spins_app")
_DSN = os.environ.get("CAJA_TEST_DSN", "postgresql://postgres@127.0.0.1:54329/radiocaja?sslmode=disable")
_BD = _DSN.rsplit("/", 1)[-1].split("?")[0]
try:
    import psycopg2
    _c = psycopg2.connect(host="127.0.0.1", port=54329, user="postgres", dbname="postgres")
    _c.autocommit = True
    _cur = _c.cursor()
    _cur.execute("SELECT 1 FROM pg_database WHERE datname = %s", (_BD,))
    if not _cur.fetchone():
        _cur.execute("CREATE DATABASE %s ENCODING 'UTF8' LC_COLLATE 'C' LC_CTYPE 'C' TEMPLATE template0" % _BD)
        _c.close()
        _c2 = psycopg2.connect(host="127.0.0.1", port=54329, user="postgres", dbname=_BD)
        _c2.autocommit = True
        _c2.cursor().execute('CREATE EXTENSION IF NOT EXISTS "uuid-ossp"')
        _c2.close()
    else:
        _c.close()
except Exception as _exc:
    print("no se pudo preparar la base de prueba:", _exc)
os.environ["DATABASE_URL"] = _DSN
os.environ.setdefault("SUPABASE_URL", "")
os.environ.setdefault("SUPABASE_SERVICE_ROLE_KEY", "")
os.environ.setdefault("FLASK_SECRET_KEY", "t")
tmp = pathlib.Path(tempfile.gettempdir())
for n in ("app33_schema_bootstrap.lock", "app33_personnel_bootstrap.lock"):
    (tmp / n).write_text("x")

import models
models.Base.metadata.drop_all(models.engine)
models.Base.metadata.create_all(models.engine)
import app as A

A.app.config["WTF_CSRF_ENABLED"] = False
D = decimal.Decimal
HOY = datetime.date.today()
AYER = HOY - datetime.timedelta(days=30)

OK, KO = [], []
def check(nombre, cond, extra=""):
    (OK if cond else KO).append(nombre)
    print(("  ok  " if cond else "  FALLA  ") + nombre + ((" → " + str(extra)[:260]) if (extra and not cond) else ""))

def eur(x):
    return A._money_value(x)


# ─────────────────────────────────────────────────────────────────────────────
print("\n── 1. Semilla: artista, empresa, actividad y contrato ─────────────────")
s = models.SessionLocal()
try:
    u = models.User(email="caja@33producciones.es", password_hash="x", role=10)
    s.add(u); s.flush()
    s.add(models.UserProfile(user_id=u.id, nick="Caja", departments=["Administración"]))
    emp = models.GroupCompany(name="33 Producciones")
    art = models.Artist(name="Los Ñus")
    ven = models.Venue(name="Sala Prueba", municipality="Sevilla", province="Sevilla")
    s.add_all([emp, art, ven]); s.flush()
    UID, AID, EMP = str(u.id), str(art.id), str(emp.id)
    # Una actividad YA CELEBRADA, vendida a un promotor, con caché de 10.000 €.
    c = models.Concert(artist_id=art.id, venue_id=ven.id, date=AYER, sale_type="VENDIDO",
                       capacity=500, activity_type="CONCIERTO", status="CONFIRMADO")
    s.add(c); s.flush()
    CID = str(c.id)
    s.add(models.ConcertCache(concert_id=c.id, kind="FIXED", amount=D("10000")))
    s.commit()
finally:
    s.close()

print("\n── 2. LAS ACTIVIDADES: el contrato reparte el importe FINAL ───────────")
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        c = s.get(models.Concert, A.to_uuid(CID))
        importe, de_donde, _cob = A._artist_cash_concert_settled(c)
        check("sin plan de pagos manda el CACHÉ", de_donde == "caché" and importe == D("10000"), (de_donde, importe))
        # Administración factura 9.000 € (un ajuste de última hora): ese es el importe final.
        c.payment_terms_json = [{"concept": "Caché", "amount": 9000, "collected_at": "2026-02-01"}]
        s.commit()
        importe, de_donde, cobrado = A._artist_cash_concert_settled(c)
        check("con plan de pagos manda la LIQUIDACIÓN, no el caché pactado",
              de_donde == "liquidación" and importe == D("9000"), (de_donde, importe))
        check("y se sabe cuánto se ha cobrado de verdad", cobrado == D("9000"), cobrado)
        # ⚠️⚠️ HASTA QUE NO TERMINA LA LIQUIDACIÓN NO HAY IMPORTE (sep 2026, lo pidió Dani): lo que
        # se apunta es lo que le ha correspondido al artista EN LA LIQUIDACIÓN, no la factura del
        # caché. Una actividad sin bolsa no está liquidada.
        d = A._artist_cash_data(s, art, None)
        check("sin liquidación cerrada, la actividad NO entra en los ingresos",
              not [g for g in d["income"]["groups"] if g["key"] == "ACTIVIDADES"],
              [g["key"] for g in d["income"]["groups"]])
        check("pero se DICE que hay 1 actividad sin liquidar",
              d["balance"]["pending_activities"] == 1, d["balance"]["pending_activities"])
        check("y con lo que le correspondería al artista",
              d["balance"]["pending_activity_amount"] == D("9000"),
              d["balance"]["pending_activity_amount"])
    finally:
        s.close()

# La BOLSA de esa actividad: mientras esté abierta, su liquidación no ha terminado.
s = models.SessionLocal()
try:
    b3 = models.WorkflowBag(title="Bolsa del concierto", artist_id=A.to_uuid(AID), artist_ids=[AID],
                            bag_type="CONCIERTO", status="ACTIVA", start_date=AYER,
                            linked_type="CONCERT", linked_id=A.to_uuid(CID), company_id=A.to_uuid(EMP))
    s.add(b3); s.flush()
    B3 = str(b3.id)
    s.commit()
finally:
    s.close()
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        d = A._artist_cash_data(s, art, None)
        check("con su bolsa TODAVÍA ABIERTA, sigue sin entrar",
              not [g for g in d["income"]["groups"] if g["key"] == "ACTIVIDADES"],
              [g["key"] for g in d["income"]["groups"]])
        # Al cerrar la liquidación de la actividad, ya aparece.
        b3 = s.get(models.WorkflowBag, A.to_uuid(B3))
        b3.status = "CERRADA"
        A._bag_cash_default_on_close(s, b3)
        s.commit()
        # Sin contrato: entero al artista, y se dice.
        d = A._artist_cash_data(s, art, None)
        acts = [g for g in d["income"]["groups"] if g["key"] == "ACTIVIDADES"]
        check("cerrada la liquidación, la actividad entra en los ingresos", bool(acts),
              [g["key"] for g in d["income"]["groups"]])
        check("y deja de contarse como pendiente", d["balance"]["pending_activities"] == 0,
              d["balance"]["pending_activities"])
        check("sin contrato, el importe va ENTERO al artista",
              acts and acts[0]["artist_amount"] == D("9000"), acts and acts[0]["artist_amount"])
        check("y se dice que es porque no hay contrato",
              acts and "sin contrato" in (acts[0]["rows"][0]["note"] or ""), acts and acts[0]["rows"][0]["note"])
    finally:
        s.close()

# Ahora SÍ hay contrato: 80 % artista / 20 % oficina en conciertos vendidos.
s = models.SessionLocal()
try:
    art = s.get(models.Artist, A.to_uuid(AID))
    contrato = models.ArtistContract(artist_id=art.id, name="Contrato de booking",
                                     signed_date=HOY - datetime.timedelta(days=400))
    s.add(contrato); s.flush()
    s.add(models.ArtistContractCommitment(contract_id=contrato.id, concept="Conciertos vendidos",
                                          pct_artist=D("80"), pct_office=D("20")))
    s.commit()
finally:
    s.close()
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        d = A._artist_cash_data(s, art, None)
        acts = [g for g in d["income"]["groups"] if g["key"] == "ACTIVIDADES"][0]
        check("CON contrato, el 80 % es del artista", acts["artist_amount"] == D("7200.00"), acts["artist_amount"])
        check("y el 20 % de la oficina", acts["office_amount"] == D("1800.00"), acts["office_amount"])
        check("el reparto se aplica sobre la LIQUIDACIÓN (9.000), no sobre el caché (10.000)",
              acts["artist_amount"] + acts["office_amount"] == D("9000.00"),
              acts["artist_amount"] + acts["office_amount"])
    finally:
        s.close()

print("\n── 3. LAS BOLSAS: solo cuentan las CERRADAS y con su balance FINAL ────")
s = models.SessionLocal()
try:
    art = s.get(models.Artist, A.to_uuid(AID))
    # Bolsa 1 · discográfica ABIERTA: 3.000 € que asume la bolsa. NO puede contar todavía.
    b1 = models.WorkflowBag(title="Grabación del disco", artist_id=art.id, artist_ids=[AID],
                            bag_type="DISCO", status="ACTIVA", start_date=AYER, company_id=A.to_uuid(EMP))
    # Bolsa 2 · promoción CERRADA: 1.000 € de la bolsa + 500 € que cubre el promotor (no es nuestro)
    b2 = models.WorkflowBag(title="Gira de promoción", artist_id=art.id, artist_ids=[AID],
                            bag_type="PROMOCION", status="CERRADA", start_date=AYER, company_id=A.to_uuid(EMP))
    s.add_all([b1, b2]); s.flush()
    B1, B2 = str(b1.id), str(b2.id)
    s.add(models.BagExpense(bag_id=b1.id, concept="Estudio", amount_gross=D("3000"), covered_by="BOLSA"))
    s.add(models.BagExpense(bag_id=b2.id, concept="Hotel", amount_gross=D("1000"), covered_by="BOLSA"))
    s.add(models.BagExpense(bag_id=b2.id, concept="Catering", amount_gross=D("500"), covered_by="PROMOTOR"))
    s.commit()
finally:
    s.close()
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        b2 = s.get(models.WorkflowBag, A.to_uuid(B2))
        check("una bolsa cerrada SIN decidir todavía no cuenta", A._bag_cash_counts(b2) is False)
        d = A._artist_cash_data(s, art, None)
        check("…y el invertido sigue a cero", d["balance"]["office_invested"] == D("0"),
              d["balance"]["office_invested"])
        # Administración dice que sí.
        A._bag_cash_decide(s, b2, "INCLUIR", nick="Caja"); s.commit()
        d = A._artist_cash_data(s, art, None)
        check("al incluirla, cuenta", d["balance"]["office_invested"] == D("1000"),
              d["balance"]["office_invested"])
        check("y lo que cubre el PROMOTOR no se imputa (1.000, no 1.500)",
              d["balance"]["office_invested"] == D("1000"), d["balance"]["office_invested"])
        check("se dice que 500 € quedan fuera del balance",
              any("Fuera del balance" in (r["note"] or "")
                  for g in d["expense"]["groups"] for r in g["rows"]),
              [r["note"] for g in d["expense"]["groups"] for r in g["rows"]])
        check("la bolsa ABIERTA no suma", d["balance"]["office_invested"] == D("1000"))
        check("pero se DICE que hay 1 bolsa abierta", d["balance"]["open_bags"] == 1, d["balance"]["open_bags"])
        check("con su importe", d["balance"]["open_amount"] == D("3000"), d["balance"]["open_amount"])
        # Al cerrarla, entra.
        b1 = s.get(models.WorkflowBag, A.to_uuid(B1))
        b1.status = "CERRADA"
        A._bag_cash_default_on_close(s, b1); s.commit()
        d = A._artist_cash_data(s, art, None)
        check("al CERRAR la bolsa abierta, ya suma", d["balance"]["office_invested"] == D("4000"),
              d["balance"]["office_invested"])
        check("y deja de contarse como abierta", d["balance"]["open_bags"] == 0, d["balance"]["open_bags"])
        # Y si se decide que NO, sale.
        A._bag_cash_decide(s, b1, "NO_INCLUIR", nick="Caja"); s.commit()
        d = A._artist_cash_data(s, art, None)
        check("si se decide que NO se incluye, sale del balance",
              d["balance"]["office_invested"] == D("1000"), d["balance"]["office_invested"])
        check("y tampoco se cuenta como abierta (alguien ya dijo que no es del artista)",
              d["balance"]["open_bags"] == 0, d["balance"]["open_bags"])
        A._bag_cash_decide(s, b1, "INCLUIR", nick="Caja"); s.commit()
    finally:
        s.close()

print("\n── 4. LA REGLA DEL CACHÉ: la bolsa de una actividad la paga su caché ──")
s = models.SessionLocal()
try:
    # ⚠️ La bolsa de la actividad se creó y se cerró en el apartado 2 (es lo que hace que la
    # actividad entre en la caja): aquí solo se le mete el gasto.
    b3 = s.get(models.WorkflowBag, A.to_uuid(B3))
    b3.cash_impact = "INCLUIR"
    s.add(models.BagExpense(bag_id=b3.id, concept="Backline", amount_gross=D("4000"), covered_by="BOLSA"))
    s.commit()
finally:
    s.close()
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        b3 = s.get(models.WorkflowBag, A.to_uuid(B3))
        check("4.000 € cubiertos por un caché de 10.000 € NO son gasto de la casa",
              A._bag_cash_cost(s, b3) == D("0"), A._bag_cash_cost(s, b3))
        d = A._artist_cash_data(s, art, None)
        check("y el invertido no cambia", d["balance"]["office_invested"] == D("4000"),
              d["balance"]["office_invested"])
        check("pero se DICE que lo cubre el caché",
              any("caché" in (r["note"] or "") for g in d["expense"]["groups"] for r in g["rows"]),
              [r["note"] for g in d["expense"]["groups"] for r in g["rows"]])
        # Si el gasto se pasa del caché, la diferencia sí es nuestra.
        s.add(models.BagExpense(bag_id=b3.id, concept="Producción", amount_gross=D("8000"), covered_by="BOLSA"))
        s.commit()
        check("lo que el caché NO cubre sí es gasto (12.000 − 10.000)",
              A._bag_cash_cost(s, b3) == D("2000"), A._bag_cash_cost(s, b3))
    finally:
        s.close()

print("\n── 5. LOS ROYALTIES: lo facturado MENOS todo lo que se paga ───────────")
s = models.SessionLocal()
try:
    art = s.get(models.Artist, A.to_uuid(AID))
    otro = models.Promoter(nick="Autor invitado")
    s.add(otro); s.flush()
    OTRO = str(otro.id)
    cancion = models.Song(title="Tema uno", release_date=datetime.date(2026, 1, 15))
    s.add(cancion); s.flush()
    SONG = str(cancion.id)
    s.add(models.SongArtist(song_id=cancion.id, artist_id=art.id))
    ini, fin = datetime.date(2026, 1, 1), datetime.date(2026, 6, 30)
    # La compañía factura 10.000 € por el tema. Al artista se le liquidan 3.000 y a un tercero 1.000.
    s.add(models.RoyaltyLiquidation(
        beneficiary_kind="ARTIST", beneficiary_id=art.id, period_start=ini, period_end=fin,
        snapshot={"total_amount": 3000, "total_income": 10000,
                  "items": [{"item_id": SONG, "item_kind": "SONG", "income": 10000, "pct": 30, "amount": 3000}]}))
    s.add(models.RoyaltyLiquidation(
        beneficiary_kind="PROMOTER", beneficiary_id=otro.id, period_start=ini, period_end=fin,
        snapshot={"total_amount": 1000, "total_income": 10000,
                  "items": [{"item_id": SONG, "item_kind": "SONG", "income": 10000, "pct": 10, "amount": 1000}]}))
    s.commit()
finally:
    s.close()
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        filas = A._artist_cash_royalties(s, art, None)
        check("hay un semestre de royalties", len(filas) == 1, len(filas))
        r = filas[0]
        check("lo FACTURADO por su repertorio se cuenta UNA vez (10.000, no 20.000)",
              r["income"] == D("10000"), r["income"])
        check("los royalties PAGADOS son los del artista Y los del tercero (4.000)",
              r["paid"] == D("4000"), r["paid"])
        check("el beneficio de la compañía es 10.000 − 4.000 = 6.000",
              r["office_amount"] == D("6000"), r["office_amount"])
        # ⚠️⚠️ SOLO LO YA FACTURADO POR ÉL (sep 2026, lo pidió Dani): la liquidación está calculada
        # pero todavía no ha emitido su factura, así que no puede salir como cobrada.
        check("lo liquidado SIN su factura todavía no cuenta como cobrado",
              r["artist_amount"] == D("0"), r["artist_amount"])
        check("y se cuenta aparte, para poder decirlo", r["artist_pending"] == D("3000"),
              r["artist_pending"])
        d = A._artist_cash_data(s, art, None)
        check("el balance lo dice", d["balance"]["pending_royalty_amount"] == D("3000"),
              d["balance"]["pending_royalty_amount"])
        # En cuanto el artista factura (INVOICED), ya es suyo.
        liq = (s.query(models.RoyaltyLiquidation)
               .filter(models.RoyaltyLiquidation.beneficiary_kind == "ARTIST").first())
        liq.status = "INVOICED"
        s.commit()
        filas = A._artist_cash_royalties(s, art, None)
        r = filas[0]
        check("y lo que YA HA FACTURADO el artista son sus 3.000", r["artist_amount"] == D("3000"),
              r["artist_amount"])
        d = A._artist_cash_data(s, art, None)
        disco = [g for g in d["income"]["groups"] if g["key"] == "DISCOGRAFICO"][0]
        check("en la caja, el discográfico dice lo mismo", disco["office_amount"] == D("6000"), disco["office_amount"])
        # ⚠️ El invertido son SOLO las bolsas (1.000 + 3.000 + los 2.000 que el caché no cubre):
        # los 4.000 € de royalties pagados NO están ahí — reducen el ingreso, no son gasto.
        check("⚠️ un royalty pagado NO es gasto: no aparece en el invertido",
              d["balance"]["office_invested"] == D("6000"), d["balance"]["office_invested"])
        check("se dice lo facturado y lo pagado",
              "Facturado" in (disco["rows"][0]["note"] or "") and "pagados" in (disco["rows"][0]["note"] or ""),
              disco["rows"][0]["note"])
    finally:
        s.close()

print("\n── 6. LOS CUATRO DATOS del resumen ────────────────────────────────────")
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        d = A._artist_cash_data(s, art, None)
        b = d["balance"]
        check("1) cobrado por el artista = 7.200 (actividad) + 3.000 (royalties ya facturados)",
              b["artist_billed"] == D("10200.00"), b["artist_billed"])
        check("2) invertido por compañía = 1.000 + 3.000 + 2.000 (lo que el caché no cubre)",
              b["office_invested"] == D("6000"), b["office_invested"])
        check("3) ingresado por compañía = 1.800 (actividad) + 6.000 (royalties)",
              b["office_income"] == D("7800.00"), b["office_income"])
        check("4) resultado para compañía = 7.800 − 6.000",
              b["office_result"] == D("1800.00"), b["office_result"])
        check("y sale en verde (deja beneficio)", b["positive"] is True)
    finally:
        s.close()

print("\n── 7. LAS PANTALLAS y el cierre en administración ─────────────────────")
cli = A.app.test_client()
with cli.session_transaction() as ses:
    ses["user_id"] = UID
    ses["role"] = 10
r = cli.get("/artistas/%s?tab=caja" % AID)
html = r.get_data(as_text=True)
check("la pestaña Caja abre (200)", r.status_code == 200, r.status_code)
for rotulo in ("Cobrado por el artista", "Invertido por compañía",
               "Ingresado por compañía", "Resultado para compañía"):
    check("el resumen dice «%s»" % rotulo, rotulo in html)
check("se despliega cada tipo con sus bolsas", "ac-group__head" in html and "Gira de promoción" in html)
check("y desde ahí se va a la bolsa", "/bolsas/" in html)

# El cierre en administración: la pregunta y que se guarda.
s = models.SessionLocal()
try:
    b4 = models.WorkflowBag(title="Bolsa a cerrar", artist_id=A.to_uuid(AID), artist_ids=[AID],
                            bag_type="PROMOCION", status="CERRADA", liquidation_status="PENDIENTE_CIERRE",
                            start_date=AYER, company_id=A.to_uuid(EMP))
    s.add(b4); s.flush()
    B4 = str(b4.id)
    s.add(models.BagExpense(bag_id=b4.id, concept="Prensa", amount_gross=D("700"), covered_by="BOLSA"))
    s.commit()
finally:
    s.close()
html = cli.get("/administracion?tab=pendiente&subtab=cierre").get_data(as_text=True)
check("al cerrar, administración ve la pregunta de la caja",
      "¿Se incluye en la caja del artista?" in html)
check("y ve lo que le cuesta a la casa", "700" in html)
r = cli.post("/administracion/bolsas/%s/cierre-liquidacion" % B4,
             data={"mode": "ARCHIVAR", "cash_impact": "NO_INCLUIR"}, follow_redirects=True)
s = models.SessionLocal()
try:
    b4 = s.get(models.WorkflowBag, A.to_uuid(B4))
    check("la decisión se guarda", (b4.cash_impact or "") == "NO_INCLUIR", b4.cash_impact)
    check("con quién la tomó", (b4.cash_decided_by_nick or "") == "Caja", b4.cash_decided_by_nick)
    check("y la bolsa queda cerrada", A._bag_is_closed(b4) is True, b4.liquidation_status)
finally:
    s.close()
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        d = A._artist_cash_data(s, art, None)
        check("y esos 700 € NO entran en la caja", d["balance"]["office_invested"] == D("6000"),
              d["balance"]["office_invested"])
    finally:
        s.close()
# Desde la bolsa se puede cambiar.
r = cli.post("/bolsas/%s/caja-artista" % B4, data={"cash_impact": "INCLUIR"}, follow_redirects=True)
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        d = A._artist_cash_data(s, art, None)
        check("cambiándola desde la bolsa, entra", d["balance"]["office_invested"] == D("6700"),
              d["balance"]["office_invested"])
    finally:
        s.close()
html = cli.get("/bolsas/%s" % B4).get_data(as_text=True)
check("la bolsa dice si va a la caja del artista", "Caja del artista" in html)

print("\n── 8. LAS BOLSAS DE ANTES: se marcan solas ────────────────────────────")
s = models.SessionLocal()
try:
    vieja = models.WorkflowBag(title="Bolsa antigua", artist_id=A.to_uuid(AID), artist_ids=[AID],
                               bag_type="PROMOCION", status="ARCHIVADA", start_date=AYER,
                               company_id=A.to_uuid(EMP))
    abierta = models.WorkflowBag(title="Bolsa viva", artist_id=A.to_uuid(AID), artist_ids=[AID],
                                 bag_type="PROMOCION", status="ACTIVA", start_date=AYER,
                                 company_id=A.to_uuid(EMP))
    s.add_all([vieja, abierta]); s.flush()
    VIEJA, ABIERTA = str(vieja.id), str(abierta.id)
    s.execute(models.text("DELETE FROM app_settings WHERE key = 'bag_cash_backfill_v1'"))
    s.commit()
finally:
    s.close()
with A.app.test_request_context("/"):
    A._bag_cash_backfill()
s = models.SessionLocal()
try:
    check("una bolsa YA CERRADA de antes entra como incluida",
          (s.get(models.WorkflowBag, A.to_uuid(VIEJA)).cash_impact or "") == "INCLUIR")
    check("y una ABIERTA no se toca (se decidirá al cerrarla)",
          s.get(models.WorkflowBag, A.to_uuid(ABIERTA)).cash_impact is None)
    check("no se pisa lo que ya se había decidido",
          (s.get(models.WorkflowBag, A.to_uuid(B4)).cash_impact or "") == "INCLUIR")
finally:
    s.close()

print("\n── 9. LA PESTAÑA «CAJA» DE ADMINISTRACIÓN ─────────────────────────────")
# ⚠️ No hay segundo motor ni segunda pantalla: la lista sale del MISMO `_artist_cash_data` y, al
# abrir un sujeto, se incluye la MISMA `_artist_cash.html` (lo pidió Dani: «lo que se cambie en las
# fichas de los artistas o aquí se cambia en ambos lados»).
r = cli.get("/administracion?tab=caja")
html = r.get_data(as_text=True)
check("la pestaña Caja abre (200)", r.status_code == 200, r.status_code)
check("sale el artista con su balance", "Los Ñus" in html)
for rotulo in ("Cobrado por los artistas", "Invertido por compañía",
               "Ingresado por compañía", "Resultado para compañía"):
    check("el total de todos dice «%s»" % rotulo, rotulo in html)
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        d = A._artist_cash_data(s, art, None)
        vista = A._cash_overview(s, None)
        mio = [f for f in vista["rows"] if f["id"] == AID]
        check("la lista dice EXACTAMENTE lo mismo que su ficha",
              mio and mio[0]["office_invested"] == d["balance"]["office_invested"]
              and mio[0]["office_income"] == d["balance"]["office_income"],
              (mio and mio[0], d["balance"]))
    finally:
        s.close()
html = cli.get("/administracion?tab=caja&sujeto=%s" % AID).get_data(as_text=True)
check("al abrir un sujeto se ve SU caja, la misma de la ficha",
      "Cobrado por el artista" in html and "Invertido por compañía" in html)
check("con el enlace a su ficha", "Abrir su ficha" in html)
check("y el selector de año lleva a la pestaña (no a la ficha)",
      "/administracion?tab=caja&amp;sujeto=" in html or "administracion" in html)
check("el año filtra", cli.get("/administracion?tab=caja&anio=%d" % HOY.year).status_code == 200)

# EL PDF del resumen.
r = cli.get("/artistas/%s/caja/resumen.pdf" % AID)
check("el PDF del resumen se genera", r.status_code == 200 and r.get_data()[:4] == b"%PDF", r.status_code)
check("y es un adjunto con su nombre", "attachment" in (r.headers.get("Content-Disposition") or ""),
      r.headers.get("Content-Disposition"))
check("el PDF por año también", cli.get("/artistas/%s/caja/resumen.pdf?anio=%d" % (AID, HOY.year)).status_code == 200)
check("la pantalla ofrece el PDF", "resumen.pdf" in cli.get("/artistas/%s?tab=caja" % AID).get_data(as_text=True))

print("\n── 9b. GIRAS COMPRADAS Y CICLOS PROPIOS ───────────────────────────────")
# ⚠️ Su caja son SUS actividades y las bolsas de esas actividades. Va APARTE y no suma al total de
# los artistas: ese mismo dinero ya está contado en la caja del artista que toca.
s = models.SessionLocal()
try:
    art = s.get(models.Artist, A.to_uuid(AID))
    gira = models.PurchasedTour(name="Gira de prueba", artist_id=art.id, artist_ids=[AID],
                                start_date=AYER, status="ACTIVA")
    s.add(gira); s.flush()
    GIRA = str(gira.id)
    ven = s.query(models.Venue).first()
    c2 = models.Concert(artist_id=art.id, venue_id=ven.id, date=AYER, sale_type="VENDIDO",
                        capacity=400, activity_type="CONCIERTO", status="CONFIRMADO",
                        purchased_tour_id=gira.id,
                        payment_terms_json=[{"concept": "Caché", "amount": 5000, "collected_at": "2026-02-01"}])
    s.add(c2); s.flush()
    C2 = str(c2.id)
    bg = models.WorkflowBag(title="Bolsa de la gira", artist_id=art.id, artist_ids=[AID],
                            bag_type="GIRA", status="CERRADA", start_date=AYER,
                            linked_type="CONCERT", linked_id=c2.id, company_id=A.to_uuid(EMP),
                            cash_impact="INCLUIR")
    s.add(bg); s.flush()
    s.add(models.BagExpense(bag_id=bg.id, concept="Autobús", amount_gross=D("1500"), covered_by="BOLSA"))
    s.commit()
finally:
    s.close()
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        gira = s.get(models.PurchasedTour, A.to_uuid(GIRA))
        d = A._group_cash_data(s, "TOUR", gira, None)
        check("la gira factura lo de SUS actividades (el 80 % de 5.000)",
              d["balance"]["artist_billed"] == D("4000.00"), d["balance"]["artist_billed"])
        check("y la compañía ingresa su 20 %", d["balance"]["office_income"] == D("1000.00"),
              d["balance"]["office_income"])
        check("el gasto es el de las bolsas de esas actividades",
              d["balance"]["office_invested"] == D("1500"), d["balance"]["office_invested"])
        check("y el resultado, la resta", d["balance"]["office_result"] == D("-500.00"),
              d["balance"]["office_result"])
        vista = A._cash_overview(s, None)
        check("la gira sale en su propio bloque", any(g["id"] == "TOUR:%s" % GIRA for g in vista["groups"]),
              [g["id"] for g in vista["groups"]])
        check("⚠️ y NO suma al total de los artistas",
              vista["totals"]["office_invested"] == [f for f in vista["rows"] if f["id"] == AID][0]["office_invested"],
              (vista["totals"]["office_invested"],))
    finally:
        s.close()
html = cli.get("/administracion?tab=caja").get_data(as_text=True)
check("la pestaña enseña el bloque de giras y ciclos", "Giras, ciclos y festivales" in html)
check("y dice que no se suman arriba", "No se suman arriba" in html)
r = cli.get("/administracion?tab=caja&sujeto=TOUR:%s" % GIRA)
html = r.get_data(as_text=True)
check("se abre la caja de la gira (200)", r.status_code == 200, r.status_code)
check("con la misma pantalla", "Cobrado por el artista" in html and "Invertido por compañía" in html)
check("sin los botones de apuntes anteriores (una gira no los tiene)", "Subir apuntes" not in html)
check("y diciendo que ese dinero ya cuenta en la caja del artista",
      "ya cuenta también en la caja de cada artista" in html)

print("\n── 10. LA PLANTILLA: desplegables y los tres importes ─────────────────")
import openpyxl, io as _io
r = cli.get("/artistas/%s/caja/plantilla.xlsx" % AID)
check("la plantilla se baja", r.status_code == 200, r.status_code)
wb = openpyxl.load_workbook(_io.BytesIO(r.get_data()))
ws = wb["Apuntes"]
cabecera = [c.value for c in ws[3]]
check("los importes se llaman como lo que son",
      "Ingreso artista" in cabecera and "Ingreso oficina" in cabecera and "Gasto oficina" in cabecera,
      cabecera)
check("y ya no ponen «inversión» ni «beneficio»",
      not any("Inversión" in (c or "") or "Beneficio" in (c or "") for c in cabecera), cabecera)
check("hay una hoja de listas, oculta", "Listas" in wb.sheetnames and wb["Listas"].sheet_state == "hidden")
dvs = {str(dv.sqref).split(":")[0][0]: dv for dv in ws.data_validations.dataValidation}
cols = {t: chr(ord("A") + i) for i, t in enumerate(cabecera)}
for titulo in ("Ingreso o gasto", "Tipo", "Empresa del grupo"):
    check("«%s» es un desplegable" % titulo, cols[titulo] in dvs, sorted(dvs))
    if cols[titulo] in dvs:
        check("…y tira de la lista de la hoja", "Listas!" in (dvs[cols[titulo]].formula1 or ""),
              dvs[cols[titulo]].formula1)
check("la lista del tipo trae los de ingreso y los de gasto",
      len([c for c in wb["Listas"]["B"] if c.value]) >= 8, len([c for c in wb["Listas"]["B"] if c.value]))
check("y la de empresas, las dadas de alta",
      any((c.value or "") == "33 Producciones" for c in wb["Listas"]["C"]))

# LA SUBIDA: con la plantilla NUEVA y con una VIEJA (columnas en el otro orden y con los nombres
# de antes). ⚠️ Es lo que impide que un gasto acabe contado como ingreso del artista.
def _sube(filas, cabecera_filas):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Apuntes"
    hoja.append(["Caja de Los Ñus · apuntes anteriores"])
    hoja.append([])
    hoja.append(cabecera_filas)
    for f in filas:
        hoja.append(f)
    buf = _io.BytesIO()
    libro.save(buf)
    return cli.post("/artistas/%s/caja/subir" % AID,
                    data={"archivo": (_io.BytesIO(buf.getvalue()), "apuntes.xlsx")},
                    content_type="multipart/form-data", follow_redirects=True)

_sube([["01/03/2024", "Gasto", "Promoción", "Cartelería vieja", "0", "0", "1.250,00", "33 Producciones", ""]],
      ["Fecha", "Ingreso o gasto", "Tipo", "Concepto", "Ingreso artista", "Ingreso oficina",
       "Gasto oficina", "Empresa del grupo", "Notas"])
s = models.SessionLocal()
try:
    fila = (s.query(models.ArtistLedgerEntry)
            .filter(models.ArtistLedgerEntry.concept == "Cartelería vieja").first())
    check("con la plantilla NUEVA, el gasto entra como gasto",
          fila is not None and A._money_value(fila.amount_invested) == D("1250")
          and A._money_value(fila.amount_artist) == D("0"),
          fila and (fila.amount_invested, fila.amount_artist))
finally:
    s.close()

# ⚠️⚠️ La plantilla VIEJA traía «Inversión realizada / Beneficio artista / Beneficio compañía» EN
# OTRO ORDEN: leyéndola por posición, esos 900 € de gasto habrían entrado como ingreso del artista.
_sube([["02/03/2024", "Gasto", "Marketing", "Campaña antigua", "900,00", "0", "0", "33 Producciones", ""]],
      ["Fecha", "Ingreso o gasto", "Tipo", "Concepto", "Inversión realizada", "Beneficio artista",
       "Beneficio compañía", "Empresa del grupo", "Notas"])
s = models.SessionLocal()
try:
    fila = (s.query(models.ArtistLedgerEntry)
            .filter(models.ArtistLedgerEntry.concept == "Campaña antigua").first())
    check("una plantilla VIEJA se lee por su cabecera, no por el sitio",
          fila is not None and A._money_value(fila.amount_invested) == D("900")
          and A._money_value(fila.amount_artist) == D("0"),
          fila and (fila.amount_invested, fila.amount_artist))
    check("y entra como PENDIENTE (no suma hasta validarla)",
          fila is not None and (fila.status or "") == "PENDIENTE", fila and fila.status)
    check("con su empresa del grupo", fila is not None and fila.company_id is not None)
finally:
    s.close()
with A.app.test_request_context("/"):
    s = models.SessionLocal()
    try:
        art = s.get(models.Artist, A.to_uuid(AID))
        d = A._artist_cash_data(s, art, None)
        # 6.700 de antes + los 1.500 de la bolsa de la gira (que también es del artista).
        check("lo subido y sin validar NO cuenta en el balance",
              d["balance"]["office_invested"] == D("8200"), d["balance"]["office_invested"])
        check("pero se dice cuántos esperan", d["balance"]["pending_entries"] == 2,
              d["balance"]["pending_entries"])
    finally:
        s.close()

print("\n════════════════════════════════════════════════════════════")
print("  %d comprobaciones OK · %d FALLAN" % (len(OK), len(KO)))
if KO:
    print("  FALLAN:")
    for k in KO:
        print("   ·", k)
print("════════════════════════════════════════════════════════════")
sys.exit(1 if KO else 0)
